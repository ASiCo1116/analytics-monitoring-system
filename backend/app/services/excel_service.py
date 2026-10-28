"""Excel export of the analytics results (.xlsx, built with openpyxl)."""

import asyncio
import io
from dataclasses import dataclass
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import Row, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.timezone import to_naive_local
from app.db.base import utc_now
from app.models import METRIC_UNITS, Metric, Record, User
from app.schemas.analytics import (
    AnalyticsParams,
    GroupStats,
    Interval,
    Source,
    Stats,
    TrendPoint,
)
from app.services import analytics_service

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
# The detail sheet holds at most this many rows (the newest ones)
DETAIL_LIMIT = 10_000
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
# Column widths in characters: table sheets, and the summary's label and
# value columns
TABLE_COLUMN_WIDTH = 18
SUMMARY_LABEL_WIDTH = 12
SUMMARY_VALUE_WIDTH = 40
SOURCE_LABELS = {Source.records: "資料記錄", Source.metrics: "即時監控資料"}
STATS_HEADER = ["筆數", "總計", "平均", "最大值", "最小值"]


@dataclass
class Report:
    params: AnalyticsParams
    start: datetime | None
    end: datetime | None
    stats: Stats
    groups: list[GroupStats]
    trend: list[TrendPoint]
    detail_header: list[str]
    detail_rows: list[list[Any]]
    exported_at: datetime


def _local(value: datetime | None) -> datetime | None:
    """Stored naive UTC to naive APP_TIMEZONE time (Excel has no zones)."""
    return None if value is None else to_naive_local(value)


async def _detail(
    db: AsyncSession, params: AnalyticsParams
) -> tuple[list[str], list[list[Any]]]:
    """Header and the newest DETAIL_LIMIT rows matching the filters."""
    if params.source == Source.records:
        header = ["ID", "標題", "數值", "分類", "記錄時間", "建立者", "描述"]
        query = (
            select(
                Record.id,
                Record.title,
                Record.value,
                Record.category,
                Record.recorded_at,
                User.username,
                Record.description,
            )
            .join(Record.owner)
            .order_by(Record.recorded_at.desc(), Record.id.desc())
        )

        def to_row(row: Row) -> list[Any]:
            return [
                row.id,
                row.title,
                row.value,
                row.category,
                _local(row.recorded_at),
                row.username,
                row.description,
            ]

    else:
        header = ["記錄時間", "伺服器", "指標", "數值", "單位", "閾值", "異常"]
        query = select(
            Metric.recorded_at,
            Metric.server,
            Metric.metric,
            Metric.value,
            Metric.threshold,
            Metric.is_anomaly,
        ).order_by(Metric.recorded_at.desc(), Metric.id.desc())

        def to_row(row: Row) -> list[Any]:
            return [
                _local(row.recorded_at),
                row.server,
                row.metric.value,
                row.value,
                METRIC_UNITS[row.metric],
                row.threshold,
                row.is_anomaly,
            ]

    query = query.where(*analytics_service.conditions(params)).limit(
        DETAIL_LIMIT
    )
    return header, [to_row(row) for row in await db.execute(query)]


async def build_export(
    db: AsyncSession, params: AnalyticsParams
) -> tuple[str, bytes]:
    """(file name, .xlsx content) for the given filters.

    Uses the same service functions as the analytics endpoints, so the
    numbers match what the dashboard shows.
    """
    # Resolve the range once: with the metrics default ("last 24 hours")
    # each sheet would otherwise take its own "now"
    start, end = params.time_range()
    params = params.model_copy(update={"start": start, "end": end})
    detail_header, detail_rows = await _detail(db, params)
    report = Report(
        params=params,
        start=start,
        end=end,
        stats=await analytics_service.summary(db, params),
        groups=await analytics_service.groups(db, params),
        trend=await analytics_service.trend(db, params, Interval.hour),
        detail_header=detail_header,
        detail_rows=detail_rows,
        exported_at=_local(utc_now().replace(microsecond=0)),
    )
    # Building the workbook is CPU work: keep it off the event loop so the
    # simulator and WebSocket pushes are not held up
    content = await asyncio.to_thread(_render, report)
    filename = (
        f"analytics_{params.source}_{report.exported_at:%Y%m%d_%H%M}.xlsx"
    )
    return filename, content


def _write_table(
    sheet: Worksheet, header: list[str], rows: list[list[Any]]
) -> None:
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = DATETIME_FORMAT
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[
            column[0].column_letter
        ].width = TABLE_COLUMN_WIDTH


def _stats_values(stats: Stats) -> list[Any]:
    return [stats.count, stats.sum, stats.avg, stats.max, stats.min]


def _render(report: Report) -> bytes:
    params = report.params
    workbook = Workbook()

    summary = workbook.active
    summary.title = "摘要"
    rows: list[list[Any]] = [
        ["資料來源", SOURCE_LABELS[params.source]],
        ["開始時間", _local(report.start) or "不限"],
        ["結束時間", _local(report.end) or "不限"],
    ]
    if params.source == Source.records:
        rows.append(["分類", params.category or "全部"])
    else:
        rows.append(["指標", params.metric.value])
        rows.append(["伺服器", params.server or "全部"])
    rows += [
        ["時區", settings.app_timezone],
        ["匯出時間", report.exported_at],
        [],
        *(
            [name, value]
            for name, value in zip(
                STATS_HEADER, _stats_values(report.stats), strict=True
            )
        ),
    ]
    if report.stats.count > DETAIL_LIMIT:
        # The statistics above still cover every matching row
        rows += [
            [],
            ["備註", f"明細已截斷，僅含最新 {DETAIL_LIMIT:,} 筆"],
        ]
    for row in rows:
        summary.append(row)
    for row in summary.iter_rows():
        row[0].font = Font(bold=True)
        if len(row) > 1 and isinstance(row[1].value, datetime):
            row[1].number_format = DATETIME_FORMAT
    summary.column_dimensions["A"].width = SUMMARY_LABEL_WIDTH
    summary.column_dimensions["B"].width = SUMMARY_VALUE_WIDTH

    group_name = "分類" if params.source == Source.records else "伺服器"
    _write_table(
        workbook.create_sheet("分組聚合"),
        [group_name, *STATS_HEADER],
        [[g.group, *_stats_values(g)] for g in report.groups],
    )
    _write_table(
        workbook.create_sheet("趨勢（依小時）"),
        ["時間", "筆數", "平均", "最大值", "最小值"],
        [
            [_local(p.bucket), p.count, p.avg, p.max, p.min]
            for p in report.trend
        ],
    )
    _write_table(
        workbook.create_sheet("明細"), report.detail_header, report.detail_rows
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
