import io
import re
from datetime import datetime, timedelta

from openpyxl import load_workbook

from app.db.base import utc_now
from app.models import Metric, MetricName, Record, Role
from app.services import excel_service
from app.services.excel_service import XLSX_MEDIA_TYPE

EXPORT = "/api/v1/analytics/export"
SHEETS = ["摘要", "分組聚合", "趨勢（依小時）", "明細"]


def summary_values(workbook) -> dict:
    return {
        row[0]: row[1] if len(row) > 1 else None
        for row in workbook["摘要"].iter_rows(values_only=True)
        if row and row[0]
    }


async def test_viewer_downloads_records(client, db, users, auth_headers):
    recorded_at = datetime(2026, 10, 1, 1, 30)  # 09:30 in Taipei
    db.add_all(
        Record(
            title=f"Sale {value}",
            value=value,
            category=category,
            recorded_at=recorded_at,
            owner_id=users[Role.user].id,
        )
        for category, value in (("sales", 10), ("sales", 20), ("traffic", 30))
    )
    await db.commit()

    response = await client.get(
        EXPORT, params={"source": "records"}, headers=auth_headers[Role.viewer]
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    disposition = response.headers["content-disposition"]
    assert re.search(
        r'filename="analytics_records_\d{8}_\d{4}\.xlsx"', disposition
    )

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == SHEETS
    summary = summary_values(workbook)
    assert summary["資料來源"] == "資料記錄"
    assert summary["筆數"] == 3
    assert summary["總計"] == 60
    assert "備註" not in summary
    groups = list(workbook["分組聚合"].iter_rows(values_only=True))
    assert groups == [
        ("分類", "筆數", "總計", "平均", "最大值", "最小值"),
        ("sales", 2, 30, 15, 20, 10),
        ("traffic", 1, 30, 30, 30, 30),
    ]
    trend = list(workbook["趨勢（依小時）"].iter_rows(values_only=True))
    # Local time, without a zone (Excel has none)
    assert trend[1][:2] == (datetime(2026, 10, 1, 9, 0), 3)
    detail = list(workbook["明細"].iter_rows(values_only=True))
    assert len(detail) == 4
    assert detail[1][4] == datetime(2026, 10, 1, 9, 30)
    assert detail[1][5] == "user"


async def test_detail_is_truncated(client, db, auth_headers, monkeypatch):
    monkeypatch.setattr(excel_service, "DETAIL_LIMIT", 5)
    newest = utc_now().replace(microsecond=0)
    db.add_all(
        Metric(
            server="server-1",
            metric=MetricName.cpu,
            value=i,
            threshold=90,
            is_anomaly=False,
            recorded_at=newest - timedelta(seconds=i),
        )
        for i in range(8)
    )
    await db.commit()

    response = await client.get(
        EXPORT,
        params={"source": "metrics", "metric": "cpu"},
        headers=auth_headers[Role.admin],
    )

    workbook = load_workbook(io.BytesIO(response.content))
    summary = summary_values(workbook)
    assert summary["筆數"] == 8  # statistics still use every row
    assert summary["指標"] == "cpu"
    assert summary["備註"] == "明細已截斷，僅含最新 5 筆"
    detail = list(workbook["明細"].iter_rows(values_only=True))
    assert len(detail) == 6  # header + 5
    assert [row[3] for row in detail[1:]] == [0, 1, 2, 3, 4]  # newest first
    assert detail[1][4] == "%"


async def test_export_validates_filters(client, auth_headers):
    response = await client.get(
        EXPORT, params={"source": "metrics"}, headers=auth_headers[Role.viewer]
    )

    assert response.status_code == 422
